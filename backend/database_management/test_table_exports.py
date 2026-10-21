import csv
import io
import json
import sqlite3
import tempfile
import zipfile
from contextlib import nullcontext
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook
from django.test import SimpleTestCase, TestCase, override_settings
from django.contrib.auth import get_user_model
from rest_framework.test import APIClient

from .table_exports import assemble, write_insert, dictionary_columns
from . import adapters, transfers
from .models import DatabaseAsset, DatabaseTransferTask


class TableExportAssemblyTests(SimpleTestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.asset = SimpleNamespace(name="sample-db", db_type="mysql")
        self.tables = []
        for index, name in enumerate(("users", "orders")):
            part = self.root / f"{index}.part"
            part.write_text('["id", "name"]\n[1, "中文"]\n', encoding="utf-8")
            self.tables.append((name, part))

    def export_format(self, fmt):
        task = SimpleNamespace(format=fmt, database="tele", schema="", asset=self.asset)
        return assemble(task, self.root, self.tables, lambda: None)

    def test_csv_writes_one_headered_file_per_table(self):
        paths = self.export_format("csv")
        self.assertEqual({path.name for path in paths}, {"users.csv", "orders.csv"})
        with paths[0].open(encoding="utf-8", newline="") as source:
            self.assertEqual(list(csv.reader(source)), [["id", "name"], ["1", "中文"]])

    def test_json_and_txt_match_database_table_layout(self):
        json_path = self.export_format("json")[0]
        data = json.loads(json_path.read_text(encoding="utf-8"))
        self.assertEqual(data["users"], [{"id": 1, "name": "中文"}])
        self.assertEqual(data["orders"], [{"id": 1, "name": "中文"}])
        txt = self.export_format("txt")[0].read_bytes().decode("utf-8")
        self.assertIn("-----------\r\nusers\r\n-----------\r\n", txt)
        self.assertIn("1\t中文\r\n", txt)

    def test_excel_has_one_sheet_per_table_and_dictionary_has_reference_columns(self):
        path = self.export_format("xlsx")[0]
        workbook = load_workbook(path, read_only=True, data_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["users", "orders"])
        self.assertEqual(list(workbook["users"].values), [("id", "name"), (1, "中文")])
        dictionary_part = self.root / "dictionary.part"
        dictionary_part.write_text(json.dumps([[1, "id", "int", "", "", "No", "", "主键", "", "", ""]], ensure_ascii=False), encoding="utf-8")
        task = SimpleNamespace(format="dictionary", database="tele", schema="", asset=self.asset)
        dictionary_path = assemble(task, self.root, [("users", dictionary_part)], lambda: None)[0]
        dictionary_book = load_workbook(dictionary_path, read_only=False, data_only=True)
        sheet = dictionary_book["Sheet1"]
        self.assertIn("A1:K1", {item.coord for item in sheet.merged_cells.ranges})
        self.assertEqual(sheet["A2"].value, "PriKey")
        self.assertEqual(sheet["B2"].value, "Name")
        self.assertEqual(sheet["B3"].value, "id")

    def test_sql_insert_options_change_batch_and_upsert_syntax(self):
        task = SimpleNamespace(asset=self.asset)
        batched = io.StringIO()
        write_insert(batched, task, "users", ["id", "name"], [(1, "A"), (2, "B")], {"batchInsert": True, "upsert": True})
        batched.seek(0); text = batched.read()
        self.assertIn("(1, 'A'),\n(2, 'B')", text)
        self.assertIn("ON DUPLICATE KEY UPDATE", text)
        separate = io.StringIO()
        write_insert(separate, task, "users", ["id"], [(1,), (2,)], {"batchInsert": False})
        separate.seek(0)
        self.assertEqual(separate.read().count("INSERT INTO"), 2)

    def test_dictionary_preserves_native_metadata_for_other_drivers(self):
        for kind in ['postgresql', 'kingbase', 'sqlserver', 'oracle', 'dameng', 'clickhouse']:
            task = SimpleNamespace(asset=SimpleNamespace(db_type=kind, username='owner', database='tele'), database='tele', schema='owner')
            with self.subTest(kind=kind), patch.object(adapters, 'columns', return_value=[{'name': 'amount', 'type': 'numeric', 'nullable': 'YES', 'default': None}]), patch.object(adapters, 'run', return_value=([{'name': 'amount', 'length': 12, 'scale': 2, 'comment': '金额'}], 1)):
                row = dictionary_columns(task, object(), 'orders')[0]
                self.assertEqual(row[3:6], [12, 2, 'Yes'])
                self.assertEqual(row[7], '金额')


class TableExportTaskTests(TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.settings_override = override_settings(DATABASE_TRANSFER_DIR=self.temp.name, DATABASE_ASSET_SQLITE_ROOT=self.temp.name,
                                                   DATABASE_TRANSFER_MIN_FREE_BYTES=0, DATABASE_TRANSFER_MAX_RESULT_BYTES=0, DATABASE_TRANSFER_BATCH_ROWS=2)
        self.settings_override.enable(); self.addCleanup(self.settings_override.disable)
        with sqlite3.connect(self.root / 'sample.db') as conn:
            conn.execute('CREATE TABLE users (id INTEGER PRIMARY KEY, name TEXT, empty TEXT)')
            conn.executemany('INSERT INTO users VALUES (?, ?, ?)', [(1, '中文', None), (2, '=1+1', ''), (3, 'comma,line\nnext', '')])
            conn.execute('CREATE TABLE empty_table (id INTEGER)')
        conn.close()
        self.user = get_user_model().objects.create_superuser(username='exporter', password='test-password')
        self.asset = DatabaseAsset.objects.create(name='sqlite', db_type='sqlite', options={'file': 'sample.db'})
        self.client = APIClient(); self.client.force_authenticate(self.user)
        self.emit_patch = patch.object(transfers, 'emit'); self.emit_patch.start(); self.addCleanup(self.emit_patch.stop)

    def task(self, fmt, **extra):
        response = self.client.post('/api/database-management/transfers/', {
            'assetId': self.asset.id, 'direction': 'export', 'scope': 'database', 'database': 'main',
            'tables': ['users', 'empty_table'], 'format': fmt, 'exportVersion': 2, **extra,
        }, format='json')
        self.assertEqual(response.status_code, 202, response.data)
        return DatabaseTransferTask.objects.get(pk=response.data['id'])

    def run_task(self, task):
        claimed = transfers.claim_one()
        self.assertEqual(claimed.id, task.id)
        transfers.execute_claimed(claimed)
        task.refresh_from_db()
        return task

    def test_five_formats_round_trip_empty_tables_unicode_and_all_rows(self):
        for fmt in ['json', 'csv', 'txt', 'xlsx', 'dictionary']:
            with self.subTest(format=fmt):
                task = self.run_task(self.task(fmt))
                self.assertEqual(task.status, 'succeeded', task.error)
                self.assertEqual(task.checkpoints['completedTables'], ['users', 'empty_table'])
                with zipfile.ZipFile(task.output_path) as archive:
                    if fmt == 'json':
                        data = json.loads(archive.read('main.json'))
                        self.assertEqual(len(data['users']), 3)
                        self.assertIsNone(data['users'][0]['empty'])
                        self.assertEqual(data['empty_table'], [])
                    elif fmt == 'csv':
                        self.assertEqual(set(archive.namelist()), {'users.csv', 'empty_table.csv'})
                        rows = list(csv.reader(io.StringIO(archive.read('users.csv').decode('utf-8'))))
                        self.assertEqual(rows[-1][1], 'comma,line\nnext')
                        self.assertEqual(archive.read('empty_table.csv'), b'id\r\n')
                    elif fmt == 'xlsx':
                        book = load_workbook(io.BytesIO(archive.read('main.xlsx')))
                        self.assertEqual(book['users']['B3'].value, '=1+1')
                        self.assertEqual(book['users']['B3'].data_type, 's')
                        self.assertEqual(book['users'].max_row, 4)
                    elif fmt == 'dictionary':
                        book = load_workbook(io.BytesIO(archive.read('main数据字典.xlsx')))
                        self.assertEqual(book['Sheet1']['B3'].value, 'id')
                    else:
                        self.assertIn('中文', archive.read('main.txt').decode('utf-8'))

    def test_single_table_exports_are_downloadable_in_the_selected_format(self):
        for fmt, extension in [('json', 'json'), ('csv', 'csv'), ('txt', 'txt'), ('xlsx', 'xlsx'), ('dictionary', 'xlsx')]:
            with self.subTest(format=fmt):
                task = self.run_task(self.task(fmt, scope='table', objectName='users', tables=['users'], fileName=f'users.{extension}'))
                self.assertEqual(task.status, 'succeeded', task.error)
                self.assertEqual(task.source_name, f'users.{extension}')
                self.assertEqual(Path(task.output_path).suffix, f'.{extension}')
                if fmt == 'json':
                    self.assertEqual(len(json.loads(Path(task.output_path).read_text(encoding='utf-8'))['users']), 3)
                elif fmt == 'csv':
                    with Path(task.output_path).open(encoding='utf-8', newline='') as source:
                        self.assertEqual(len(list(csv.reader(source))), 4)
                elif fmt == 'xlsx':
                    book = load_workbook(task.output_path, read_only=True, data_only=True)
                    self.addCleanup(book.close)
                    self.assertEqual(len(list(book['users'].values)), 4)
                elif fmt == 'dictionary':
                    book = load_workbook(task.output_path, read_only=True, data_only=True)
                    self.addCleanup(book.close)
                    self.assertEqual(book['Sheet1']['B3'].value, 'id')
                else:
                    self.assertIn('id\tname', Path(task.output_path).read_text(encoding='utf-8'))

    def test_cancel_and_permission_revocation_stop_export(self):
        task = self.task('json')
        task.cancel_requested = True; task.save(update_fields=['cancel_requested'])
        self.run_task(task)
        self.assertEqual(task.status, 'cancelled')
        other = self.task('csv')
        with patch.object(transfers, 'require_task_permissions', return_value=False): self.run_task(other)
        self.assertEqual(other.status, 'failed')

    def test_continue_on_error_omits_failed_table_and_reports_partial_result(self):
        original = adapters.columns
        def columns(asset, database, table, schema):
            if table == 'users': raise RuntimeError('driver error containing private data')
            return original(asset, database, table, schema)
        task = self.task('json', options={'continueOnError': True})
        with patch.object(adapters, 'columns', side_effect=columns): self.run_task(task)
        self.assertEqual(task.status, 'succeeded', task.error)
        self.assertEqual(task.checkpoints['failedTables'], ['users'])
        self.assertNotIn('private data', str(task.checkpoints))
        with zipfile.ZipFile(task.output_path) as archive:
            self.assertEqual(json.loads(archive.read('main.json')), {'empty_table': []})
        stop = self.task('json', options={'continueOnError': False})
        with patch.object(adapters, 'columns', side_effect=columns): self.run_task(stop)
        self.assertEqual(stop.status, 'failed')

    def test_limit_cannot_be_bypassed_by_continue_on_error_and_download_is_owner_only(self):
        task = self.task('csv', options={'continueOnError': True})
        with override_settings(DATABASE_TRANSFER_MAX_ROWS=1): self.run_task(task)
        self.assertEqual(task.status, 'failed')
        good = self.run_task(self.task('json'))
        other = get_user_model().objects.create_superuser(username='other', password='test-password')
        self.client.force_authenticate(other)
        self.assertEqual(self.client.get(f'/api/database-management/transfers/{good.id}/download/').status_code, 404)

    def test_excel_invalid_cell_skips_only_failed_table_when_requested(self):
        conn = sqlite3.connect(self.root / 'sample.db')
        try:
            conn.execute('UPDATE users SET name=? WHERE id=1', ('x' * 32768,))
            conn.commit()
        finally: conn.close()
        task = self.run_task(self.task('xlsx', options={'continueOnError': True}))
        self.assertEqual(task.status, 'succeeded', task.error)
        self.assertEqual(task.checkpoints['failedTables'], ['users'])
        with zipfile.ZipFile(task.output_path) as archive:
            book = load_workbook(io.BytesIO(archive.read('main.xlsx')))
            self.assertEqual(book.sheetnames, ['empty_table'])
            book.close()
        stop = self.run_task(self.task('xlsx', options={'continueOnError': False}))
        self.assertEqual(stop.status, 'failed')

    def test_legacy_sql_named_zip_downloads_as_real_zip(self):
        result = self.root / 'legacy.result.sql'
        result.write_text('CREATE TABLE sample (id INT);\n', encoding='utf-8')
        task = transfers.create_task(owner=self.user, asset=self.asset, direction='export', scope='database',
                                                  format='sql', database='tele', status='succeeded', source_name='tele-sql.zip',
                                                  output_path=str(result), parameters={'content': 'structure_data'})
        response = self.client.get(f'/api/database-management/transfers/{task.id}/download/')
        self.assertEqual(response.status_code, 200)
        try: body = b''.join(response.streaming_content)
        finally: response.close()
        with zipfile.ZipFile(io.BytesIO(body)) as archive:
            self.assertEqual(archive.namelist(), ['tele.sql'])
            self.assertEqual(archive.read('tele.sql'), result.read_bytes())
        self.assertEqual(response['Content-Type'], 'application/zip')
        self.assertEqual(result.suffix, '.sql')

    def test_sql_bundle_honors_structure_data_drop_and_upsert_options(self):
        from . import mysql_dump, table_exports
        self.asset.db_type = 'mysql'; self.asset.save(update_fields=['db_type'])
        for content in ['structure_data', 'data', 'structure']:
            with self.subTest(content=content):
                task = self.task('sql', tables=['users'], content=content, options={'dropTable': True, 'batchInsert': False, 'upsert': True})
                with patch.object(adapters, 'object_list', return_value=[{'name': 'users'}]), patch.object(adapters, 'connection', side_effect=lambda *args: nullcontext(object())), patch.object(table_exports, 'table_columns', return_value=[{'COLUMN_NAME': 'id', 'EXTRA': ''}]), patch.object(table_exports, 'row_batches', side_effect=lambda *args: nullcontext(iter([[(0,), (2,)]]))), patch.object(mysql_dump, 'query', return_value=[{'Create Table': 'CREATE TABLE `users` (`id` int)'}]) as query:
                    self.run_task(task)
                self.assertEqual(task.status, 'succeeded', task.error)
                self.assertTrue(any("NO_AUTO_VALUE_ON_ZERO" in call.args[1] for call in query.call_args_list))
                with zipfile.ZipFile(task.output_path) as archive: sql = archive.read('main.sql').decode('utf-8')
                self.assertEqual('DROP TABLE IF EXISTS' in sql, content != 'data')
                self.assertEqual('CREATE TABLE' in sql, content != 'data')
                self.assertEqual(sql.count('INSERT INTO'), 0 if content == 'structure' else 2)
                self.assertEqual('ON DUPLICATE KEY UPDATE' in sql, content != 'structure')


