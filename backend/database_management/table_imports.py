import csv
import json
import re
import zipfile
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.conf import settings
from . import adapters
from .mysql_dump import DumpError, quote, resource_guard, generated_column


FORMATS = {"xlsx", "csv", "json", "txt"}


def read_rows(task):
    path = Path(task.input_path)
    if task.format == "json":
        if path.stat().st_size > 64 * 1024 * 1024: raise DumpError("JSON 文件超过 64 MiB，请拆分或改用 CSV")
        with path.open(encoding="utf-8-sig") as stream: payload = json.load(stream)
        if isinstance(payload, dict):
            if set(payload) != {task.object_name}: raise DumpError("JSON 必须只包含当前表名对应的数据数组")
            payload = payload[task.object_name]
        if not isinstance(payload, list) or not payload: raise DumpError("JSON 必须是非空对象数组")
        if not isinstance(payload[0], dict): raise DumpError("JSON 数据行必须为对象")
        fields = list(payload[0])
        yield fields
        for row in payload:
            if not isinstance(row, dict) or set(row) != set(fields): raise DumpError("JSON 每行必须包含相同字段")
            yield [row[field] for field in fields]
        return
    if task.format == "xlsx":
        from openpyxl import load_workbook
        with path.open("rb") as stream:
            if not zipfile.is_zipfile(stream): raise DumpError("Excel 文件格式无效")
            stream.seek(0)
            book = load_workbook(stream, read_only=True, data_only=False)
            try:
                rows = book.worksheets[0].iter_rows()
                header = next(rows, ())
                yield [str(cell.value or "") for cell in header]
                for cells in rows:
                    if any(cell.data_type == "f" for cell in cells): raise DumpError("Excel 含公式，请先转换为值")
                    values = [cell.value for cell in cells]
                    if any(value is not None for value in values): yield values
            finally: book.close()
        return
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream, delimiter="\t" if task.format == "txt" else ",")
        yield next(reader, [])
        for row in reader:
            if row: yield row


def target_columns(task):
    quote(task.object_name)
    columns = adapters.columns(task.asset, task.database, task.object_name, task.schema)
    if not columns: raise DumpError("非 SQL 导入只允许写入当前已存在的表")
    with adapters.connection(task.asset, task.database) as conn:
        rows, _ = adapters.run(conn, "SELECT COLUMN_NAME AS name, EXTRA AS extra FROM information_schema.COLUMNS WHERE TABLE_SCHEMA=%s AND TABLE_NAME=%s", (task.database, task.object_name), db_type=task.asset.db_type)
        engines, _ = adapters.run(conn, "SELECT ENGINE AS engine FROM information_schema.TABLES WHERE TABLE_SCHEMA=%s AND TABLE_NAME=%s AND TABLE_TYPE='BASE TABLE'", (task.database, task.object_name), db_type=task.asset.db_type)
        if not engines or str(engines[0]["engine"]).lower() != "innodb": raise DumpError("数据导入要求目标为 InnoDB 表，以保证失败时回滚")
    extras = {row["name"]: row["extra"] for row in rows}
    return [{**column, "extra": extras.get(column["name"], "")} for column in columns]


def mapping_for(fields, columns, mapping):
    if not fields or any(not isinstance(field, str) or not field for field in fields) or len(set(fields)) != len(fields): raise DumpError("首行字段名为空或重复")
    if not isinstance(mapping, dict) or set(mapping) - set(fields): raise DumpError("字段映射无效")
    names = {column["name"]: column for column in columns}
    targets = [mapping.get(field, field) for field in fields]
    if any(not isinstance(name, str) or name not in names for name in targets): raise DumpError("文件包含未映射的目标字段，请核对字段映射")
    if len(set(targets)) != len(targets): raise DumpError("多个来源字段不能映射到同一目标字段")
    if any(generated_column(names[name].get("extra", "")) for name in targets): raise DumpError("生成字段不能导入")
    missing = [column["name"] for column in columns if column["name"] not in targets and column["nullable"] == "NO" and column.get("default") is None and "auto_increment" not in column.get("extra", "").lower() and not generated_column(column.get("extra", ""))]
    if missing: raise DumpError("缺少必填字段：" + "、".join(missing))
    return targets, [names[name] for name in targets]


def typed_value(value, column):
    if value is None:
        if column["nullable"] == "NO" and "auto_increment" not in column.get("extra", "").lower(): raise ValueError()
        return None
    kind = column["type"].lower()
    if re.match(r"(?:tinyint|smallint|mediumint|bigint|int|integer)\b", kind):
        number = Decimal(str(value))
        if number != number.to_integral_value(): raise ValueError()
        bits = next((size for name, size in [("tinyint", 8), ("smallint", 16), ("mediumint", 24), ("bigint", 64)] if kind.startswith(name)), 32)
        low, high = (0, 2 ** bits - 1) if "unsigned" in kind else (-(2 ** (bits - 1)), 2 ** (bits - 1) - 1)
        if number < low or number > high: raise ValueError()
        return int(number)
    if re.match(r"(?:decimal|numeric|float|double|real)\b", kind):
        number = Decimal(str(value))
        if not number.is_finite(): raise ValueError()
        return number
    if kind.startswith(("datetime", "timestamp")): return value if isinstance(value, datetime) else datetime.fromisoformat(str(value))
    if kind == "date": return value if isinstance(value, date) else date.fromisoformat(str(value))
    if kind.startswith("time"): return value if isinstance(value, time) else time.fromisoformat(str(value))
    if kind == "json":
        text = json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else str(value)
        json.loads(text)
        return text
    if isinstance(value, (list, dict)): raise ValueError()
    text = str(value)
    limit = re.match(r"(?:var)?char\((\d+)\)", kind)
    if limit and len(text) > int(limit[1]): raise ValueError()
    return text


def validate_row(row, columns, number):
    if len(row) != len(columns): raise DumpError(f"第 {number} 行字段数与首行不一致")
    values = []
    for value, column in zip(row, columns):
        try: values.append(typed_value(value, column))
        except (ValueError, TypeError, InvalidOperation): raise DumpError(f"第 {number} 行字段 {column['name']} 类型或长度不匹配")
    return values


def inspect_rows(task, report, mapping=None, strict=True):
    columns = target_columns(task)
    iterator = read_rows(task)
    count, validation_error, targets, fields = 0, "", [], []
    try:
        fields = next(iterator)
        try: targets, mapped = mapping_for(fields, columns, mapping if mapping is not None else task.parameters.get("fieldMapping", {}))
        except DumpError as exc:
            if strict: raise
            validation_error = str(exc)
        for count, row in enumerate(iterator, 1):
            if not validation_error:
                try: validate_row(row, mapped, count + 1)
                except DumpError as exc:
                    if strict: raise
                    validation_error = str(exc)
            if settings.DATABASE_TRANSFER_MAX_ROWS and count > settings.DATABASE_TRANSFER_MAX_ROWS: raise DumpError("行数超过系统配置上限")
            if count % 500 == 0: resource_guard(task); report(task, stage="validating_rows", rows=count)
    finally:
        close = getattr(iterator, "close", None)
        if close: close()
    if not count: raise DumpError("文件没有可导入的数据行")
    return {"format": task.format, "columns": fields, "targetColumns": [column["name"] for column in columns],
            "fieldMapping": dict(zip(fields, targets)), "validationError": validation_error, "rows": count, "requiresExplicitConfirmation": True,
            "objects": [{"name": task.object_name, "exists": True, "hasStructure": False, "rows": count, "appendCompatible": True}],
            "warnings": ["仅导入当前表数据。覆盖会删除目标表的全部数据；跳过策略跳过主键或唯一键冲突的行。"]}


def import_rows(task, report):
    mapping = task.parameters.get("fieldMapping", {})
    preview = inspect_rows(task, report, mapping)
    columns = target_columns(task)
    iterator = read_rows(task)
    fields, mapped = mapping_for(next(iterator), columns, mapping)
    sql = f"INSERT INTO {quote(task.object_name)} ({', '.join(quote(field) for field in fields)}) VALUES ({', '.join(['%s'] * len(fields))})"
    with adapters.connection(task.asset, task.database) as conn:
        try:
            adapters.run(conn, "START TRANSACTION", db_type=task.asset.db_type)
            if task.conflict_policy == "overwrite": adapters.run(conn, f"DELETE FROM {quote(task.object_name)}", db_type=task.asset.db_type)
            for count, row in enumerate(iterator, 1):
                try: adapters.run(conn, sql, tuple(validate_row(row, mapped, count + 1)), db_type=task.asset.db_type)
                except Exception as exc:
                    if task.conflict_policy != "skip" or not exc.args or exc.args[0] != 1062: raise
                if count % 250 == 0: resource_guard(task); report(task, stage="importing_rows", rows=count, total=preview["rows"])
            report(task, stage="committing", rows=preview["rows"], total=preview["rows"])
            conn.commit()
        except BaseException:
            conn.rollback()
            raise
        finally:
            close = getattr(iterator, "close", None)
            if close: close()
