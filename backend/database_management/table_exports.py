from .credentials import resolve_credentials
"""Streaming table exports. Temporary per-table files keep partial failures isolated."""

import csv
import json
import re
import shutil
import tempfile
import zipfile
from contextlib import contextmanager
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path

from django.conf import settings
from django.utils import timezone

from . import adapters
from .mysql_dump import DumpError, generated_column, resource_guard, sql_literal, table_columns

FORMATS = {"sql", "xlsx", "dictionary", "txt", "csv", "json"}
DICTIONARY_HEADERS = ["PriKey", "Name", "FullType", "Length", "Scala", "Nullable", "Default", "Comment", "Charset", "Collation", "Extra"]


class ExportLimitError(DumpError):
    """A global resource limit must stop the task even when errors may be skipped."""


def safe_name(value):
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value)).strip(" .")[:120] or "database"
    if re.fullmatch(r"(?i)(con|prn|aux|nul|com[1-9]|lpt[1-9])(?:\..*)?", value): value = "_" + value
    return value


def scalar(value):
    if value is None or isinstance(value, (str, bool, int)): return value
    if isinstance(value, float):
        import math
        return value if math.isfinite(value) else str(value)
    if isinstance(value, (bytes, bytearray, memoryview)): return "0x" + bytes(value).hex()
    if isinstance(value, datetime): return value.isoformat(sep=" ")
    if isinstance(value, (date, time)): return value.isoformat()
    if isinstance(value, (Decimal, timedelta)): return str(value)
    if isinstance(value, dict): return {str(k): scalar(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)): return [scalar(v) for v in value]
    if hasattr(value, "read"): return scalar(value.read())
    return str(value)


def text_value(value):
    value = scalar(value)
    if value is None: return ""
    return json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else str(value)


def quote_identifier(name, kind):
    if not name or "\x00" in str(name): raise DumpError("数据库对象名称不合法")
    name = str(name)
    if kind in {"mysql", "mariadb", "clickhouse"}: return "`" + name.replace("`", "``") + "`"
    if kind == "sqlserver": return "[" + name.replace("]", "]]") + "]"
    return '"' + name.replace('"', '""') + '"'


def qualified(task, table):
    kind = task.asset.db_type
    owner = task.schema or ("public" if kind in {"postgresql", "kingbase"} else "dbo" if kind == "sqlserver" else resolve_credentials(task.asset)[0].upper() if kind in {"oracle", "dameng"} else "")
    return (quote_identifier(owner, kind) + "." if owner else "") + quote_identifier(table, kind)


@contextmanager
def row_batches(task, conn, table, fields):
    kind = task.asset.db_type
    sql = f"SELECT {', '.join(quote_identifier(f, kind) for f in fields)} FROM {qualified(task, table)}"
    batch_size = max(1, settings.DATABASE_TRANSFER_BATCH_ROWS)
    if kind == "clickhouse":
        with conn.query_row_block_stream(sql, settings={"max_result_rows": 0, "max_block_size": batch_size}) as stream:
            yield iter(stream)
        return
    if kind in {"mysql", "mariadb"}:
        from pymysql.cursors import SSCursor
        cursor = conn.cursor(SSCursor)
    elif kind in {"postgresql", "kingbase"}:
        from psycopg.rows import tuple_row
        cursor = conn.cursor(name="captain_export", row_factory=tuple_row)
    else:
        cursor = conn.cursor()
    try:
        cursor.execute(sql)
        def batches():
            while True:
                rows = cursor.fetchmany(batch_size)
                if not rows: return
                yield rows
        yield batches()
    finally:
        cursor.close()


def dictionary_columns(task, conn, table):
    if task.asset.db_type in {"mysql", "mariadb"}:
        from .mysql_dump import query
        return [[1 if c["COLUMN_KEY"] == "PRI" else "", c["COLUMN_NAME"], c["COLUMN_TYPE"],
                 c.get("CHARACTER_MAXIMUM_LENGTH") or c.get("NUMERIC_PRECISION") or "", c.get("NUMERIC_SCALE") if c.get("NUMERIC_SCALE") is not None else "",
                 "Yes" if c["IS_NULLABLE"] == "YES" else "No", text_value(c.get("COLUMN_DEFAULT")), c.get("COLUMN_COMMENT", ""),
                 c.get("CHARACTER_SET_NAME") or "", c.get("COLLATION_NAME") or "", c.get("EXTRA") or ""]
                for c in query(conn, "SELECT * FROM information_schema.columns WHERE TABLE_SCHEMA=%s AND TABLE_NAME=%s ORDER BY ORDINAL_POSITION", (task.database, table))]
    columns = adapters.columns(task.asset, task.database, table, task.schema)
    kind = task.asset.db_type
    owner = task.schema or ("dbo" if kind == "sqlserver" else "public" if kind in {"postgresql", "kingbase"} else resolve_credentials(task.asset)[0].upper())
    details = []
    if kind in {"postgresql", "kingbase"}:
        sql = """SELECT c.column_name AS name, COALESCE(c.character_maximum_length, c.numeric_precision) AS length,
            c.numeric_scale AS scale, c.collation_name AS collation, c.character_set_name AS charset,
            pg_catalog.format_type(a.atttypid, a.atttypmod) AS full_type,
            pg_catalog.col_description(t.oid, a.attnum) AS comment, c.is_identity AS extra
            FROM information_schema.columns c JOIN pg_catalog.pg_namespace n ON n.nspname=c.table_schema
            JOIN pg_catalog.pg_class t ON t.relnamespace=n.oid AND t.relname=c.table_name
            JOIN pg_catalog.pg_attribute a ON a.attrelid=t.oid AND a.attname=c.column_name
            WHERE c.table_schema=%s AND c.table_name=%s"""
        details, _ = adapters.run(conn, sql, (owner, table), db_type=kind, max_rows=max(500, len(columns)))
    elif kind == "sqlserver":
        sql = """SELECT c.COLUMN_NAME AS name, COALESCE(c.CHARACTER_MAXIMUM_LENGTH, c.NUMERIC_PRECISION) AS length,
            c.NUMERIC_SCALE AS scale, c.COLLATION_NAME AS collation, c.CHARACTER_SET_NAME AS charset,
            CAST(p.value AS nvarchar(4000)) AS comment,
            CASE WHEN sc.is_identity=1 THEN 'IDENTITY' WHEN sc.is_computed=1 THEN 'COMPUTED' ELSE '' END AS extra
            FROM INFORMATION_SCHEMA.COLUMNS c JOIN sys.schemas s ON s.name=c.TABLE_SCHEMA
            JOIN sys.tables t ON t.schema_id=s.schema_id AND t.name=c.TABLE_NAME
            JOIN sys.columns sc ON sc.object_id=t.object_id AND sc.name=c.COLUMN_NAME
            LEFT JOIN sys.extended_properties p ON p.major_id=t.object_id AND p.minor_id=sc.column_id AND p.name='MS_Description' AND p.class=1
            WHERE c.TABLE_SCHEMA=%s AND c.TABLE_NAME=%s"""
        details, _ = adapters.run(conn, sql, (owner, table), db_type=kind, max_rows=max(500, len(columns)))
    elif kind in {"oracle", "dameng"}:
        sql = f"""SELECT c.COLUMN_NAME AS name, COALESCE(c.DATA_PRECISION, c.CHAR_LENGTH, c.DATA_LENGTH) AS length,
            c.DATA_SCALE AS scale, m.COMMENTS AS comment FROM ALL_TAB_COLUMNS c
            LEFT JOIN ALL_COL_COMMENTS m ON m.OWNER=c.OWNER AND m.TABLE_NAME=c.TABLE_NAME AND m.COLUMN_NAME=c.COLUMN_NAME
            WHERE c.OWNER={adapters.marker(kind, 1)} AND c.TABLE_NAME={adapters.marker(kind, 2)}"""
        details, _ = adapters.run(conn, sql, (owner, table), db_type=kind, max_rows=max(500, len(columns)))
    elif kind == "clickhouse":
        details, _ = adapters.run(conn, "SELECT name, comment, default_kind AS extra FROM system.columns WHERE database={db:String} AND table={table:String}",
                                  {"db": task.database or task.asset.database or "default", "table": table}, db_type=kind, max_rows=max(500, len(columns)))
    extra_columns = {c["name"]: c for c in adapters.metadata_rows(details)}
    for column in columns:
        extra = extra_columns.get(column["name"], {})
        column.update({key: value if value is not None else "" for key, value in extra.items() if key != "name"})
        if extra.get("full_type"): column["type"] = extra["full_type"]
    return [[1 if c.get("column_key") == "PRI" else "", c["name"], c["type"], c.get("length", ""), c.get("scale", ""),
             "Yes" if str(c.get("nullable")).upper() in {"YES", "Y"} else "No", text_value(c.get("default")), c.get("comment", ""),
             c.get("charset", ""), c.get("collation", ""), c.get("extra", "")] for c in columns]


def validate_xlsx_values(values):
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    for value in values:
        value = text_value(value) if isinstance(value, (dict, list)) else scalar(value)
        if isinstance(value, str) and (len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value)):
            raise DumpError("字段内容超出 Excel 文本限制，请改用 JSON/TXT")


def xlsx_cells(sheet, values, header=False):
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    cells = []
    for value in values:
        value = scalar(value)
        if isinstance(value, (dict, list)): value = text_value(value)
        if isinstance(value, int) and abs(value) >= 10 ** 15: value = str(value)
        if isinstance(value, str) and len(value) > 32767: raise DumpError("单元格超过 Excel 的 32767 字符限制，请改用 JSON/TXT")
        cell = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str): cell.data_type = "s"  # Database strings must never become formulas.
        if header:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8EEF5")
        cells.append(cell)
    return cells


def export_bundle(task, output_path, report):
    from .transfers import require_task_permissions, update
    fmt, options = task.format, task.parameters.get("options", {})
    tables = task.parameters.get("tables") or [item["name"] for item in adapters.object_list(task.asset, task.database, task.schema, "table")]
    available = {item["name"] for item in adapters.object_list(task.asset, task.database, task.schema, "table")}
    if not tables or any(name not in available for name in tables): raise DumpError("请选择存在的表进行导出")
    if settings.DATABASE_TRANSFER_MAX_TABLES and len(tables) > settings.DATABASE_TRANSFER_MAX_TABLES: raise DumpError("表数量超过系统配置上限")
    checkpoints = {"completedTables": [], "failedTables": [], "currentTable": "", "tableRows": 0, "tableTotal": None, "completedSteps": 0, "totalSteps": len(tables), "logs": []}
    total_rows = 0
    def log(message, level="info"):
        checkpoints["logs"] = [*checkpoints["logs"], {"time": timezone.now().isoformat(), "level": level, "message": message}][-300:]
    def heartbeat(stage, byte_count=0):
        if not require_task_permissions(task.owner, task): raise PermissionError("任务所需权限已被撤销")
        try: resource_guard(task, byte_count)
        except DumpError as exc: raise ExportLimitError(str(exc)) from exc
        report(task, stage=stage, rows=total_rows, byte_count=byte_count)
        update(task, checkpoints=checkpoints.copy(), progress=min(99, int(checkpoints["completedSteps"] * 100 / len(tables))))
    with tempfile.TemporaryDirectory(prefix=f"{task.pk}-export-", dir=output_path.parent) as directory:
        root = Path(directory)
        completed = []
        with adapters.connection(task.asset, task.database) as conn:
            if task.asset.db_type in {"mysql", "mariadb"}:
                from .mysql_dump import query
                from .mysql_connection import prepare_export_snapshot
                prepare_export_snapshot(conn)
            for index, table in enumerate(tables):
                checkpoints.update(currentTable=table, tableRows=0, tableTotal=None)
                log(f"开始导出：{table}")
                part = root / f"{index}.part"
                heartbeat(f"exporting_table:{table}")
                try:
                    if fmt == "dictionary":
                        data = dictionary_columns(task, conn, table)
                        validate_xlsx_values([table])
                        for row in data: validate_xlsx_values(row)
                        part.write_text(json.dumps(data, ensure_ascii=False, default=str), encoding="utf-8")
                    else:
                        if fmt == "sql":
                            columns = [{"name": c["COLUMN_NAME"]} for c in table_columns(conn, task.database, table) if not generated_column(c["EXTRA"])]
                        else:
                            columns = adapters.columns(task.asset, task.database, table, task.schema)
                        fields = [c["name"] for c in columns]
                        if not fields: raise DumpError("表没有可导出的字段")
                        if fmt == "xlsx":
                            if len(fields) > 16384: raise DumpError("字段数量超出 Excel 限制")
                            validate_xlsx_values(fields)
                        content = task.parameters.get("content", "data")
                        with part.open("w", encoding="utf-8", newline="") as output:
                            if fmt == "sql" and content != "data":
                                from .mysql_dump import query, quote
                                ddl = query(conn, f"SHOW CREATE TABLE {quote(table)}")[0]["Create Table"]
                                output.write(f"-- {table} DDL\n")
                                if options.get("dropTable"): output.write(f"DROP TABLE IF EXISTS {quote(table)};\n")
                                output.write(ddl.rstrip(";") + ";\n")
                            if fmt != "sql": output.write(json.dumps(fields, ensure_ascii=False) + "\n")
                            if fmt != "sql" or content != "structure":
                                with row_batches(task, conn, table, fields) as batches:
                                    for rows in batches:
                                        if fmt == "sql":
                                            write_insert(output, task, table, fields, rows, options)
                                        else:
                                            for row in rows:
                                                values = [row[f] for f in fields] if isinstance(row, dict) else list(row)
                                                if fmt == "xlsx": validate_xlsx_values(values)
                                                output.write(json.dumps([scalar(v) for v in values], ensure_ascii=False) + "\n")
                                        total_rows += len(rows)
                                        checkpoints["tableRows"] += len(rows)
                                        if fmt == "xlsx" and checkpoints["tableRows"] > 1048575: raise DumpError("表超过 Excel 行数限制，请改用 CSV/JSON")
                                        if settings.DATABASE_TRANSFER_MAX_ROWS and total_rows > settings.DATABASE_TRANSFER_MAX_ROWS: raise ExportLimitError("行数超过系统配置上限")
                                        output.flush()
                                        heartbeat(f"exporting_table:{table}", sum(p.stat().st_size for p in root.iterdir() if p.is_file()))
                        checkpoints["tableTotal"] = checkpoints["tableRows"]
                    completed.append((table, part))
                    checkpoints["completedTables"].append(table)
                    log(f"导出完成：{table}", "success")
                except (InterruptedError, PermissionError, ExportLimitError):
                    raise
                except Exception:
                    part.unlink(missing_ok=True)
                    checkpoints["failedTables"].append(table)
                    log(f"导出失败：{table}，请检查连接、读取权限和字段类型", "error")
                    update(task, checkpoints=checkpoints.copy())
                    if not options.get("continueOnError"): raise
                    if task.asset.db_type in {"postgresql", "kingbase"}: conn.rollback()
                checkpoints["completedSteps"] = index + 1
                heartbeat(f"exporting_table:{table}")
        if not completed: raise DumpError("所有表均导出失败")
        log("正在生成文件")
        heartbeat("packaging")
        files = assemble(task, root, completed, lambda: heartbeat("packaging", sum(p.stat().st_size for p in root.iterdir() if p.is_file())))
        if task.scope == "table":
            if len(files) != 1: raise DumpError("单表导出结果数量无效")
            shutil.copyfile(files[0], output_path)
            heartbeat("packaging", output_path.stat().st_size)
        else:
            with zipfile.ZipFile(output_path, "w", compression=zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
                for path in files:
                    heartbeat("packaging", output_path.stat().st_size)
                    with path.open("rb") as source, archive.open(path.name, "w", force_zip64=True) as destination:
                        while chunk := source.read(1024 * 1024):
                            destination.write(chunk)
                            heartbeat("packaging", output_path.stat().st_size)
        log("部分表导出失败，请查看结果" if checkpoints["failedTables"] else "全部导出完成", "warning" if checkpoints["failedTables"] else "success")
        update(task, checkpoints=checkpoints.copy(), preview={**task.preview, "files": [{"name": p.name, "size": p.stat().st_size} for p in files], "warnings": [f"未导出：{t}" for t in checkpoints["failedTables"]]})


def write_insert(output, task, table, fields, rows, options):
    from .mysql_dump import quote
    prefix = f"INSERT INTO {quote(table)} ({', '.join(quote(f) for f in fields)}) VALUES "
    suffix = ""
    if options.get("upsert"):
        suffix = " ON DUPLICATE KEY UPDATE " + ", ".join(f"{quote(f)}=VALUES({quote(f)})" for f in fields)
    values = ["(" + ", ".join(sql_literal(v) for v in ([r[f] for f in fields] if isinstance(r, dict) else r)) + ")" for r in rows]
    if options.get("batchInsert", True):
        # Keep generated INSERT statements consumable by the existing SQL importer.
        limit = settings.DATABASE_TRANSFER_MAX_SQL_STATEMENT_BYTES
        batch, size = [], len((prefix + suffix).encode("utf-8"))
        for value in values:
            value_size = len(value.encode("utf-8")) + 2
            if size + value_size > limit and batch:
                output.write(prefix + ",\n".join(batch) + suffix + ";\n")
                batch, size = [], len((prefix + suffix).encode("utf-8"))
            if size + value_size > limit: raise DumpError("单行 SQL 超过语句长度限制")
            batch.append(value); size += value_size
        if batch: output.write(prefix + ",\n".join(batch) + suffix + ";\n")
    else:
        for value in values: output.write(prefix + value + suffix + ";\n")


def assemble(task, root, tables, heartbeat):
    fmt = task.format
    base = safe_name(tables[0][0] if getattr(task, "scope", "") == "table" else task.database or task.schema or task.asset.name)
    if fmt in {"xlsx", "dictionary"}:
        return [assemble_workbook(task, root, tables, heartbeat)]
    paths, used = [], set()
    output = None
    if fmt != "csv":
        path = root / f"{base}.{fmt}"
        paths.append(path); output = path.open("w", encoding="utf-8", newline="")
    try:
        if fmt == "json": output.write("{\n")
        if fmt == "sql": output.write("SET FOREIGN_KEY_CHECKS = 0;\nSET NAMES utf8mb4;\nSET SQL_MODE = 'NO_AUTO_VALUE_ON_ZERO';\n")
        json_index = 0
        for index, (table, part) in enumerate(tables):
            heartbeat()
            if fmt == "sql":
                with part.open(encoding="utf-8") as source:
                    while chunk := source.read(1024 * 1024):
                        output.write(chunk)
                        output.flush()
                        heartbeat()
                continue
            if fmt == "csv":
                name = safe_name(table)
                candidate, count = name, 1
                while candidate.casefold() in used:
                    count += 1; candidate = f"{name}_{count}"
                used.add(candidate.casefold())
                path = root / f"{candidate}.csv"; paths.append(path)
                output = path.open("w", encoding="utf-8", newline="")
            with part.open(encoding="utf-8") as source:
                fields = json.loads(next(source))
                if fmt == "csv": writer = csv.writer(output, lineterminator="\r\n"); writer.writerow(fields)
                if fmt == "txt":
                    if getattr(task, "scope", "") == "table":
                        writer = csv.writer(output, delimiter="\t", lineterminator="\r\n"); writer.writerow(fields)
                    else: output.write(f"-----------\r\n{table}\r\n-----------\r\n")
                if fmt == "json": output.write((",\n" if json_index else "") + json.dumps(table, ensure_ascii=False) + ": [")
                for row_index, line in enumerate(source):
                    values = json.loads(line)
                    if fmt == "json": output.write((",\n" if row_index else "") + json.dumps(dict(zip(fields, values)), ensure_ascii=False))
                    elif fmt == "csv": writer.writerow([text_value(v) for v in values])
                    elif fmt == "txt" and getattr(task, "scope", "") == "table": writer.writerow([text_value(v) for v in values])
                    else: output.write("\t".join(text_value(v).replace("\\", "\\\\").replace("\t", "\\t").replace("\r", "\\r").replace("\n", "\\n") for v in values) + "\r\n")
                    if row_index % 500 == 0: output.flush(); heartbeat()
                if fmt == "json": output.write("]"); json_index += 1
            if fmt == "csv": output.close(); output = None
        if fmt == "json": output.write("\n}\n")
        if fmt == "sql": output.write("SET FOREIGN_KEY_CHECKS = 1;\n")
    finally:
        if output: output.close()
    return paths


def assemble_workbook(task, root, tables, heartbeat):
    from openpyxl import Workbook
    dictionary = task.format == "dictionary"
    workbook = Workbook(write_only=not dictionary)
    sheet = workbook.active if dictionary else None
    if sheet: sheet.title = "Sheet1"
    used = set()
    dictionary_row = 1
    for table, part in tables:
        heartbeat()
        if dictionary:
            sheet.merged_cells.add(f"A{dictionary_row}:K{dictionary_row}")
            sheet.append(xlsx_cells(sheet, [table], header=True))
            sheet.append(xlsx_cells(sheet, DICTIONARY_HEADERS, header=True))
            rows = json.loads(part.read_text(encoding="utf-8"))
            for row in rows: sheet.append(xlsx_cells(sheet, row))
            sheet.append([])
            dictionary_row += len(rows) + 3
            continue
        name = re.sub(r"[\\/*?:\[\]]", "_", table).strip("'")[:31] or "Table"
        candidate, count = name, 1
        while candidate.casefold() in used:
            count += 1; candidate = name[:31-len(str(count))-1] + "_" + str(count)
        used.add(candidate.casefold())
        sheet = workbook.create_sheet(candidate)
        sheet.freeze_panes = "A2"
        with part.open(encoding="utf-8") as source:
            sheet.append(xlsx_cells(sheet, json.loads(next(source)), header=True))
            for index, line in enumerate(source):
                if index >= 1048575: raise DumpError("表超过 Excel 行数限制，请改用 CSV/JSON")
                sheet.append(xlsx_cells(sheet, json.loads(line)))
                if index % 500 == 0: heartbeat()
    path = root / (safe_name(task.database or task.schema or task.asset.name) + ("数据字典" if dictionary else "") + ".xlsx")
    workbook.save(path)
    return path
